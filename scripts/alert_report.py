#!/usr/bin/env python3
"""
SOC daily alert report.

Reads alerts exported from a SIEM (CSV), builds an Excel report with a summary
by severity, technique and host, and optionally emails it.

Usage:
    python alert_report.py alerts.csv                 # build report only
    python alert_report.py alerts.csv --email         # build and email

Email settings come from environment variables (never hard-code credentials):
    SMTP_HOST, SMTP_PORT (default 587), SMTP_USER, SMTP_PASSWORD, REPORT_TO
"""
import argparse
import csv
import os
import smtplib
import sys
from collections import Counter
from datetime import date
from email.message import EmailMessage

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED = ["time", "host", "rule", "technique", "severity"]
SEVERITY_ORDER = ["critical", "high", "medium", "low", "info"]
SEVERITY_COLORS = {"critical": "C00000", "high": "FF6B6B", "medium": "FFC000", "low": "92D050", "info": "BDD7EE"}


def load_alerts(path):
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        missing = [c for c in REQUIRED if c not in (reader.fieldnames or [])]
        if missing:
            sys.exit(f"CSV is missing columns: {', '.join(missing)}")
        return [{k: (v or "").strip() for k, v in row.items()} for row in reader]


def autosize(ws):
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, 60)


def header(ws, cols):
    ws.append(cols)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")


def build_report(alerts, out_path):
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append([f"SOC Alert Report - {date.today().isoformat()}"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Total alerts: {len(alerts)}"])
    summary.append([])

    sev = Counter(a["severity"].lower() for a in alerts)
    header(summary, ["Severity", "Count"])
    for s in SEVERITY_ORDER:
        if sev.get(s):
            summary.append([s.capitalize(), sev[s]])
            summary.cell(summary.max_row, 1).fill = PatternFill("solid", fgColor=SEVERITY_COLORS[s])
    summary.append([])

    for title, key in (("MITRE technique", "technique"), ("Host", "host")):
        header(summary, [title, "Count"])
        for value, count in Counter(a[key] for a in alerts).most_common(10):
            summary.append([value, count])
        summary.append([])
    autosize(summary)

    detail = wb.create_sheet("Alerts")
    cols = list(alerts[0].keys()) if alerts else REQUIRED
    header(detail, cols)
    rank = {s: i for i, s in enumerate(SEVERITY_ORDER)}
    for a in sorted(alerts, key=lambda a: (rank.get(a["severity"].lower(), 99), a["time"])):
        detail.append([a.get(c, "") for c in cols])
    detail.freeze_panes = "A2"
    autosize(detail)

    wb.save(out_path)
    return sev


def send_email(report_path, sev, total):
    needed = ["SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "REPORT_TO"]
    missing = [v for v in needed if not os.environ.get(v)]
    if missing:
        sys.exit(f"Missing environment variables: {', '.join(missing)}")

    msg = EmailMessage()
    urgent = sev.get("critical", 0) + sev.get("high", 0)
    msg["Subject"] = f"[SOC] Daily report {date.today()} - {total} alerts ({urgent} high/critical)"
    msg["From"] = os.environ["SMTP_USER"]
    msg["To"] = os.environ["REPORT_TO"]
    lines = [f"{s.capitalize()}: {sev[s]}" for s in SEVERITY_ORDER if sev.get(s)]
    msg.set_content("Daily SOC alert summary:\n\n" + "\n".join(lines) + "\n\nFull report attached.")
    with open(report_path, "rb") as f:
        msg.add_attachment(f.read(), maintype="application",
                           subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           filename=os.path.basename(report_path))

    with smtplib.SMTP(os.environ["SMTP_HOST"], int(os.environ.get("SMTP_PORT", 587))) as s:
        s.starttls()
        s.login(os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"])
        s.send_message(msg)


def main():
    p = argparse.ArgumentParser(description="Build a daily SOC alert report from a SIEM CSV export.")
    p.add_argument("csv", help="CSV with columns: time, host, rule, technique, severity")
    p.add_argument("-o", "--output", default=f"soc_report_{date.today().isoformat()}.xlsx")
    p.add_argument("--email", action="store_true", help="Email the report (uses SMTP_* env vars)")
    args = p.parse_args()

    alerts = load_alerts(args.csv)
    sev = build_report(alerts, args.output)
    print(f"Report saved: {args.output} ({len(alerts)} alerts)")
    if args.email:
        send_email(args.output, sev, len(alerts))
        print(f"Report emailed to {os.environ['REPORT_TO']}")


if __name__ == "__main__":
    main()
